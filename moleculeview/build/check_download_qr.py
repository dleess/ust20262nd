"""Decode download QRs and verify clickable store links in the built PPTX/PDF."""
from pathlib import Path
import json

import cv2
import fitz
import numpy as np
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
slides = json.loads((ROOT / 'build/content_all.json').read_text())
ppt = Presentation(ROOT / 'build/biomolecule_view_3h_ko.pptx')
pdf = fitz.open(ROOT / 'build/biomolecule_view_3h_ko_pdb_cif.pdf')
assert len(ppt.slides) == len(pdf) == len(slides)
expected = {
    'qr_molviewapp_appstore.png': 'https://apps.apple.com/app/id6786979305',
    'qr_molapp_googleplay.png': 'https://play.google.com/store/apps/details?id=com.donghan.molapp',
}
downloads = [s for s in slides if s['kind'] == 'app_download']
assert {s['image'] for s in downloads} == set(expected)
detector = cv2.QRCodeDetector()
for s in downloads:
    url = expected[s['image']]
    assert s['download_url'] == url
    original = cv2.imread(str(ROOT / 'assets/figures' / s['image']))
    assert detector.detectAndDecode(original)[0] == url
    page = pdf[s['id'] - 1]
    for scale in (1, 2):
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pixels = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3)
        assert detector.detectAndDecode(pixels)[0] == url, (s['id'], scale)
    assert url in [link.get('uri') for link in page.get_links()]
    shapes = ppt.slides[s['id'] - 1].shapes
    qr = next(shape for shape in shapes if shape.name == s['image'])
    assert qr.click_action.hyperlink.address == url
    assert url in [run.hyperlink.address for shape in shapes if shape.has_text_frame
                   for para in shape.text_frame.paragraphs for run in para.runs]
    print(f"PASS: slide {s['id']} QR source/PDF 1x/2x and PPTX/PDF links: {url}")

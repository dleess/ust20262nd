"""Editable Korean lecture slides from the reviewed lesson content."""
from pathlib import Path
import json
from PIL import Image, ImageChops, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
FONT = 'Arial Unicode MS'
FONT_FILE = '/Library/Fonts/Arial Unicode.ttf'
NAVY = '183342'
TEAL = '277F87'
MUTED = '5A6A72'
ORANGE = 'BF570C'
WHITE = 'FFFFFF'
SOURCES = json.loads((ROOT / 'content/sources.json').read_text())
DATA = json.loads((ROOT / 'content/foundations.json').read_text()) + json.loads((ROOT / 'content/structure_reading.json').read_text())
assert [s['id'] for s in DATA] == list(range(1, 53))
assert sum(s['minutes'] for s in DATA) == 180
assert sum(s['minutes'] for s in DATA if s['kind'] == 'break') == 20
PRS = Presentation()
PRS.slide_width = Inches(13.333333)
PRS.slide_height = Inches(7.5)
PRS.core_properties.title = '생체분자 구조 읽기: 약물 설계를 위한 첫걸음'
PRS.core_properties.subject = '생물학 배경 석사생을 위한 한국어 3시간 강의'
PRS.core_properties.author = 'Donghan Lee'
PRS.core_properties.keywords = 'biomolecule, protein, ligand, EGFR, 구조 읽기'
PRS.core_properties.language = 'ko-KR'
checks = []

def color(h):
    return RGBColor.from_string(h)

def run_style(run, size, bold=False, col=NAVY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    rpr = run._r.get_or_add_rPr()
    rpr.set('lang', 'ko-KR')
    for tag in ['a:ea', 'a:cs']:
        e = OxmlElement(tag)
        e.set('typeface', FONT)
        rpr.append(e)

def text(slide, value, x, y, w, h, size=23, bold=False, col=NAVY, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = line
        run_style(r, size, bold, col)
    return shape

def estimated_lines(value, width, size):
    font = ImageFont.truetype(FONT_FILE, round(size * 2))
    limit = width * 72 * 2 - 8
    total = 0
    for raw in value.split('\n'):
        cur = ''
        total += 1
        for ch in raw:
            if font.getlength(cur + ch) > limit:
                total += 1
                cur = ch
            else:
                cur += ch
    return total

def body(slide, lines, x, y, w, h, size=23, col=NAVY):
    heights = [(estimated_lines(s, w, size) * size * 1.15 + 5) / 72 for s in lines]
    gap = 0.28
    need = sum(heights) + gap * max(0, len(lines) - 1)
    if need > h:
        size = 21
        heights = [(estimated_lines(s, w, size) * size * 1.15 + 5) / 72 for s in lines]
        need = sum(heights) + gap * max(0, len(lines) - 1)
    assert need <= h + 0.1, f'Text too dense: {lines} need {need}, have {h}'
    shape = text(slide, '\n'.join(lines), x, y, w, h, size=size, col=col)
    for p in shape.text_frame.paragraphs[:-1]:
        p.space_after = Pt(18)

def picture(slide, filename, x, y, w, h):
    path = ROOT / 'assets/figures' / filename
    im = Image.open(path).convert('RGB')
    iw, ih = im.size
    box = (0, 0, iw, ih)
    if filename.startswith(('2ITY', '1M17')) or filename in ['peptide_fragment.png','methionine_residue.png','secondary_egfr.png']:
        diff = ImageChops.difference(im, Image.new('RGB', im.size, (255,255,255))).convert('L')
        b = diff.point(lambda z: 255 if z > 20 else 0).getbbox()
        if b:
            box = (max(0,b[0]-30),max(0,b[1]-30),min(iw,b[2]+30),min(ih,b[3]+30))
    bw, bh = box[2]-box[0], box[3]-box[1]
    scale = min(w / bw, h / bh)
    pw, ph = bw * scale, bh * scale
    pic = slide.shapes.add_picture(str(path), Inches(x+(w-pw)/2), Inches(y+(h-ph)/2), width=Inches(pw), height=Inches(ph))
    pic.crop_left = box[0]/iw
    pic.crop_top = box[1]/ih
    pic.crop_right = (iw-box[2])/iw
    pic.crop_bottom = (ih-box[3])/ih
    pic.name = filename
    # Retain descriptive alt text alongside the unmodified source image.
    pic._element.nvPicPr.cNvPr.set('descr', filename.replace('_',' ').replace('.png',''))
    return pic

def table(slide, info, x=0.65, y=1.6, w=12.0, h=4.1):
    headers = info['headers']
    rows = info['rows']
    shp = slide.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h))
    tbl = shp.table
    if len(headers) == 2:
        tbl.columns[0].width = Inches(w*0.40)
        tbl.columns[1].width = Inches(w*0.60)
    for i, row in enumerate([headers] + rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i,j)
            cell.text = str(val)
            cell.margin_left = Inches(.18)
            cell.margin_right = Inches(.16)
            cell.margin_top = Inches(.08)
            cell.margin_bottom = Inches(.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(TEAL if i==0 else ('F2F6F6' if i%2 else WHITE))
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.08
                for r in p.runs:
                    run_style(r,18 if i==0 else 20,i==0,WHITE if i==0 else NAVY)
    return shp

elapsed = 0
for s in DATA:
    slide = PRS.slides.add_slide(PRS.slide_layouts[6])
    kind = s['kind']
    isdark = kind == 'break'
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY if isdark else WHITE)
    if kind == 'cover':
        text(slide, 'BIOMOLECULE VIEW', .7,.55,6,.5,17,True,TEAL)
        text(slide, '생체분자\n구조 읽기', .7,1.55,6,2.1,48,True)
        text(slide, '약물 설계를 위한 첫걸음', .7,4.0,6, .6,24)
        text(slide, 'AI 단백질 구조 · MD · docking의 기초\n한국어 강의 3시간', .7,5.05,5.7,1.05,18,col=MUTED)
        picture(slide, '2ITY_whole.png', 6.6,1.05,6.25,5.65)
    elif isdark:
        text(slide, s['title'], .8,2.0,11.8,1.05,48,True,WHITE)
        body(slide, s['body'], .8,3.55,11,1.7,24,WHITE)
    else:
        text(slide, s['title'], .65,.42,12.05,.85,32,True)
        if kind == 'image':
            picture(slide,s['image'],.55,1.45,7.5,4.95)
            body(slide,s.get('body',[]),8.35,1.8,4.25,4.65,22)
        elif kind == 'compare':
            items = s['images']
            n = len(items)
            gap=.3
            w=(12.05-gap*(n-1))/n
            for i,item in enumerate(items):
                xx=.65+i*(w+gap)
                picture(slide,item['file'],xx,1.5,w,3.85)
                text(slide,item['label'],xx,5.4,w,.7,18,True,col=TEAL)
            body(slide,s.get('body',[]),.65,6.05,12.0,.85,18)
        elif kind == 'table':
            rowcount=len(s['table']['rows'])+1
            ht=min(4.15,.76*rowcount)
            table(slide,s['table'],h=ht)
            if s.get('body'):
                body(slide,s['body'][-1:],.65,1.6+ht+.30,12.0,6.85-(1.6+ht+.30),21)
        elif kind == 'formula':
            text(slide,s['formula'],.8,1.7,11.8,1.1,36,True,col=TEAL)
            body(slide,s.get('body',[]),.85,3.15,11.65,3.35,25)
        elif kind == 'references':
            for i,key in enumerate(s['sources']):
                src=SOURCES[key]
                y=1.48+i*.73
                tx=text(slide,src['title'],.7,y,11.9,.38,17,True)
                link=text(slide,src['url'],.7,y+.34,11.9,.32,10,col=TEAL)
                for p in link.text_frame.paragraphs:
                    for r in p.runs:
                        r.hyperlink.address=src['url']
        else:
            if kind == 'activity':
                text(slide,'종이에 적거나, 짝에게 근거를 설명합니다.',.8,1.55,11.7,.5,18,col=TEAL)
            body(slide,s.get('body',[]),.8,2.45 if kind=='activity' else 1.85,11.65,4.2 if kind=='activity' else 4.75,28 if kind=='activity' else 27)
    imgs = [s.get('image','')] + [i['file'] for i in s.get('images',[])]
    external = next((f for f in imgs if f in ['water.png','peptide.png','secondary.png']),None)
    if external:
        label='OpenStax Biology 2e (Clark, Douglas & Choi), CC BY-NC-SA 4.0. 원그림 그대로 사용.'
        text(slide,label,.65,6.78,11.65,.23,9,col=MUTED)
        text(slide,'Access for free at https://openstax.org/books/biology-2e/pages/1-introduction',.65,7.01,11.6,.23,9,col=MUTED)
    elif not isdark and s['id']<51:
        if any(f.startswith(('2ITY','1M17')) for f in imgs):
            code='1M17' if any(f.startswith('1M17') for f in imgs) else '2ITY'
            text(slide,f'PDB {code} 원좌표 · 단백질 탄소 청록 / 약물 탄소 주황 / N 파랑 / O 빨강',.65,6.93,11.4,.29,11,col=MUTED)
        elif s.get('sources'):
            text(slide,'자료: '+', '.join(s['sources'])+'  ·  상세 출처는 발표자 노트와 마지막 참고 자료에 수록',.65,6.97,11.45,.25,10,col=MUTED)
    text(slide,str(s['id']),12.05,7.00,.65,.3,11,col=WHITE if isdark else MUTED,align=PP_ALIGN.RIGHT)
    end=elapsed+s['minutes']
    note=f"권장 진행: {elapsed}–{end}분 ({s['minutes']}분)\n\n{s.get('notes','')}\n\n"
    if kind=='table' and len(s.get('body',[]))>1:
        note+='추가 설명\n'+'\n'.join(s['body'][:-1])+'\n\n'
    if s.get('sources'):
        note+='출처\n'+'\n'.join(SOURCES[k]['title']+'\n'+SOURCES[k]['url']+('\n'+SOURCES[k]['detail'] if SOURCES[k].get('detail') else '') for k in s['sources'])
    if external:
        note+='\n\n그림 출처: OpenStax Biology 2e. Clark, Douglas & Choi. CC BY-NC-SA 4.0.\nhttps://creativecommons.org/licenses/by-nc-sa/4.0/\n그림을 변형하지 않고 비율을 유지해 배치했습니다.'
    slide.notes_slide.notes_text_frame.text=note
    elapsed=end
    for shape in slide.shapes:
        assert shape.left>=0 and shape.top>=0
        assert shape.left+shape.width<=PRS.slide_width+10
        assert shape.top+shape.height<=PRS.slide_height+10
    checks.append({'slide':s['id'],'title':s['title'],'minutes':s['minutes'],'objects':len(slide.shapes)})
out=ROOT/'build/biomolecule_view_3h_ko.pptx'
PRS.save(out)
(ROOT/'build/content_all.json').write_text(json.dumps(DATA,ensure_ascii=False,indent=2))
(ROOT/'build/deck_checks.json').write_text(json.dumps({'slides':52,'minutes':elapsed,'checks':checks},ensure_ascii=False,indent=2))
guide=['# 생체분자 구조 읽기: 강사 노트','총180분, 휴식20분 포함. 발표자 노트는 PPT에도 포함되어 있습니다.','']
for s in DATA:
    guide += [f"## {s['id']}. {s['title']} ({s['minutes']}분)",s.get('notes',''),'']
(ROOT/'content/teacher_notes.md').write_text('\n'.join(guide))
print(out)

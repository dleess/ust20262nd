"""Check timings, editable deck, rendered text, worksheet, and protected lesson 1."""
from pathlib import Path
import hashlib
import json
import unicodedata
import pymupdf as fitz
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
build=ROOT/'build'
data=json.loads((build/'content_all.json').read_text())
assert len(data)==62 and [s['id'] for s in data]==list(range(1,63))
assert sum(s['minutes'] for s in data)==180
assert sum(s['minutes'] for s in data if s['kind']=='break')==20
prs=Presentation(build/'msa_structure_prediction_3h_ko.pptx')
pdf=fitz.open(build/'msa_structure_prediction_3h_ko.pdf')
worksheet=fitz.open(build/'msa_structure_prediction_worksheet_ko.pdf')
assert len(prs.slides)==len(pdf)==62 and len(worksheet)==4

def norm(s):
    return ''.join(unicodedata.normalize('NFKC',s).split()).replace('\u200b','')

missing=[]
for i,(slide,page) in enumerate(zip(prs.slides,pdf),1):
    assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text
    rendered=norm(page.get_text())
    assert '\ufffd' not in rendered
    source=data[i-1]
    expected=[] if source['kind']=='references' else source.get('body',[])[:]
    if source.get('table'):
        expected += source['table']['headers'] + [str(v) for row in source['table']['rows'] for v in row]
    if source.get('sequence'):
        expected += source['sequence']['text'].splitlines()
    for value in expected:
        assert not norm(value) or norm(value) in rendered, (i,'source content missing',value)
    for shp in slide.shapes:
        assert shp.left>=0 and shp.top>=0
        assert shp.left+shp.width<=prs.slide_width+10
        assert shp.top+shp.height<=prs.slide_height+10
        chunks=[]
        if shp.has_text_frame:chunks=[p.text for p in shp.text_frame.paragraphs]
        if shp.has_table:chunks=[c.text for row in shp.table.rows for c in row.cells]
        for chunk in chunks:
            if norm(chunk) and norm(chunk) not in rendered:missing.append({'slide':i,'text':chunk})
    for block in page.get_text('blocks'):
        assert block[0]>=-1 and block[1]>=-1 and block[2]<=page.rect.width+1 and block[3]<=page.rect.height+1,(i,block[:4])
for i,page in enumerate(worksheet,1):
    assert '\ufffd' not in page.get_text()
    for block in page.get_text('blocks'):
        assert block[0]>=-1 and block[1]>=-1 and block[2]<=page.rect.width+1 and block[3]<=page.rect.height+1,('worksheet',i,block[:4])
original=json.loads((build/'lecture1_baseline_sha256.json').read_text())
for path,digest in original.items():
    assert hashlib.sha256((ROOT.parent/path).read_bytes()).hexdigest()==digest,path
query=''.join((ROOT/'examples/calmodulin_P0DP23.fasta').read_text().splitlines()[1:])
assert len(query)==149
report={'slides':62,'lecture_pdf_pages':62,'worksheet_pages':4,'total_minutes':180,'break_minutes':20,'all_slide_notes_present':True,'lecture1_outputs_unchanged':True,'missing_rendered_text':missing,'actual_structure_inference_run':False,'synthetic_practice_data':True}
(build/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
assert not missing,missing
print(json.dumps(report,ensure_ascii=False,indent=2))

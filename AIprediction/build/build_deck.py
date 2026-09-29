"""Build editable Korean lecture 2. Source content lives in content/*.json."""
from pathlib import Path
from datetime import datetime
import json
import shutil
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / 'build'
FONT = 'Arial Unicode MS'
FONT_FILE = '/Library/Fonts/Arial Unicode.ttf'
MONO_FILE = '/System/Library/Fonts/Supplemental/Courier New.ttf'
INK, TEAL, MUTED, PAPER = '183342', '277F87', '5A6A72', 'FAFAF6'
W, H = 13.333333, 7.5

def archive(path):
    if path.exists():
        folder = BUILD / 'archive' / datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        folder.mkdir(parents=True)
        shutil.copy2(path, folder / path.name)

def color(value):
    return RGBColor.from_string(value)

def lines(value, width, size, mono=False):
    # ponytail: font-width estimate; exported PowerPoint pages remain the final fit check.
    font = ImageFont.truetype(MONO_FILE if mono else FONT_FILE, round(size * 2))
    limit = (width * 72 - 6) * 2
    count = 0
    for raw in str(value).split('\n'):
        count += 1
        current = ''
        for char in raw:
            if font.getlength(current + char) > limit:
                count += 1
                current = char
            else:
                current += char
    return count

def style(run, size, bold=False, col=INK, mono=False):
    run.font.name = 'Courier New' if mono else FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    pr = run._r.get_or_add_rPr()
    pr.set('lang', 'ko-KR')
    node = OxmlElement('a:ea')
    node.set('typeface', FONT)
    pr.append(node)

def text(slide, value, x, y, w, h, size=22, bold=False, col=INK, mono=False):
    estimate = lines(value, w, size, mono) * size * 1.13 + 4
    assert estimate <= h * 72 + 5, (slide_no, str(value)[:80], estimate, h * 72)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = not mono
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    for i, raw in enumerate(str(value).split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.line_spacing = Pt(size * 1.13) if mono else 1.13
        r = p.add_run()
        r.text = raw
        style(r, size, bold, col, mono)
    return shape

def paragraphs(slide, values, x, y, w, h, size=24):
    if isinstance(values, str):
        values = [values]
    heights = [(lines(t, w, size) * size * 1.13 + 4) / 72 for t in values]
    gap = .21
    assert sum(heights) + gap * max(0, len(heights)-1) <= h + .04, (slide_no, 'body too tall',values)
    for val, height in zip(values, heights):
        text(slide, val, x, y, w, height, size)
        y += height + gap
    return y

def table(slide, info, y=1.55, size=19, x=.67, width=12, compact=False):
    matrix = [info['headers']] + info['rows']
    cols = len(matrix[0])
    widths = info.get('widths', [width / cols] * cols)
    assert len(widths) == cols and abs(sum(widths)-width)<.01
    row_heights = [max(.44 if compact else .60, max(lines(str(v), widths[j]-.30,size) for j,v in enumerate(row))*size*1.14/72+.17) for row in matrix]
    if sum(row_heights)>4.6 and size>17:
        return table(slide,info,y,size-1,x,width,compact)
    assert sum(row_heights)<=4.75, (slide_no,'table too tall',row_heights)
    shp = slide.shapes.add_table(len(matrix),cols,Inches(x),Inches(y),Inches(width),Inches(sum(row_heights)))
    tbl = shp.table
    for c, width in zip(tbl.columns,widths):
        c.width = Inches(width)
    for i, row in enumerate(matrix):
        assert len(row)==cols
        tbl.rows[i].height=Inches(row_heights[i])
        for j, value in enumerate(row):
            cell=tbl.cell(i,j)
            cell.text=str(value)
            cell.margin_left=cell.margin_right=Inches(.12)
            cell.margin_top=cell.margin_bottom=Inches(.055)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            fill_color=info.get('cell_colors',{}).get(f'{i},{j}',TEAL if i==0 else ('EEF3F2' if i%2 else 'FFFFFF'))
            cell.fill.fore_color.rgb=color(fill_color)
            for p in cell.text_frame.paragraphs:
                p.line_spacing=1.10
                for run in p.runs:
                    font_color=info.get('text_colors',{}).get(f'{i},{j}','FFFFFF' if i==0 else INK)
                    style(run,size,i==0,font_color)
    return y+sum(row_heights)

def confidence_profile(slide, values):
    chart_data=CategoryChartData()
    chart_data.categories=[str(i) for i in range(1,len(values)+1)]
    chart_data.add_series('pLDDT',values)
    chart=slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(.7),Inches(1.55),Inches(12),Inches(3.5),chart_data).chart
    chart.has_legend=False
    chart.has_title=False
    chart.font.name=FONT
    chart.font.size=Pt(16)
    for axis,label in [(chart.category_axis,'residue number'),(chart.value_axis,'pLDDT (0–100)')]:
        axis.has_title=True
        axis.axis_title.text_frame.text=label
        for paragraph in axis.axis_title.text_frame.paragraphs:
            for run in paragraph.runs:style(run,16)
        axis.tick_labels.font.name=FONT
        axis.tick_labels.font.size=Pt(15)
    chart.value_axis.minimum_scale=0
    chart.value_axis.maximum_scale=100
    chart.value_axis.major_unit=20
    chart.value_axis.has_major_gridlines=True
    chart.value_axis.major_gridlines.format.line.color.rgb=color('D8E1E1')
    chart.value_axis.major_gridlines.format.line.width=Pt(.5)
    series=chart.series[0]
    series.format.line.color.rgb=color(TEAL)
    series.format.line.width=Pt(2.2)
    series.marker.style=XL_MARKER_STYLE.CIRCLE
    series.marker.size=7
    series.marker.format.fill.solid()
    series.marker.format.fill.fore_color.rgb=color(TEAL)
    series.marker.format.line.color.rgb=color(TEAL)
    chart.plots[0].has_data_labels=True
    labels=chart.plots[0].data_labels
    labels.position=XL_LABEL_POSITION.ABOVE
    labels.font.name=FONT
    labels.font.size=Pt(14)
    labels.font.color.rgb=color(INK)

def coevolution_diagram(slide, info):
    orange = 'C57B2A'

    def stroke(points, col='A4B1B7', width=4, dashed=False):
        for (x1,y1),(x2,y2) in zip(points,points[1:]):
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
            line.line.color.rgb=color(col)
            line.line.width=Pt(width)
            if dashed:line.line.dash_style=MSO_LINE_DASH_STYLE.DASH

    def node(x,y,label,col,diameter=.58):
        shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-diameter/2),Inches(y-diameter/2),Inches(diameter),Inches(diameter))
        shape.fill.solid();shape.fill.fore_color.rgb=color(col)
        shape.line.fill.background()
        shape.text_frame.clear()
        shape.text_frame.margin_left=shape.text_frame.margin_right=0
        shape.text_frame.margin_top=shape.text_frame.margin_bottom=0
        shape.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        p=shape.text_frame.paragraphs[0];p.alignment=PP_ALIGN.CENTER
        run=p.add_run();run.text=label;style(run,21,True,'FFFFFF')

    if info['kind']=='msa_to_structure':
        text(slide,info['left_title'],.8,1.63,4.8,.45,22,True,TEAL)
        text(slide,info['right_title'],8.35,1.63,4.25,.45,22,True,TEAL)
        text(slide,'i = 20',2.55,2.27,1.1,.35,17,True,TEAL)
        text(slide,'j = 80',4.13,2.27,1.1,.35,17,True,orange)
        for x,fill in [(2.5,'DEEFEC'),(4.08,'F5E8D5')]:
            rect=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(2.7),Inches(1.0),Inches(2.03))
            rect.fill.solid();rect.fill.fore_color.rgb=color(fill);rect.line.fill.background()
        for n,(name,a,b) in enumerate(info['rows']):
            y=2.8+n*.46
            text(slide,name,.85,y,1.65,.36,18,mono=True)
            text(slide,a,2.85,y,.45,.36,21,True,TEAL,True)
            text(slide,'···',3.58,y,.42,.36,18,col=MUTED)
            text(slide,b,4.43,y,.45,.36,21,True,orange,True)
        arrow=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(5.65),Inches(3.08),Inches(1.85),Inches(.45))
        arrow.fill.solid();arrow.fill.fore_color.rgb=color(TEAL);arrow.line.fill.background()
        text(slide,info['inference'],5.43,3.78,2.5,1.0,19,True)
        # Conceptual chain path, not molecular coordinates or a predicted structure.
        path=[(8.7,2.92),(9.1,3.02),(9.5,3.25),(9.85,2.75),(10.65,2.55),(11.65,2.85),(12.05,3.5),(11.88,4.3),(11.1,4.72),(10.2,4.82),(9.42,4.5),(9.0,3.95),(9.5,3.78),(10.4,3.83),(10.85,4.1)]
        stroke(path)
        stroke([(9.5,3.25),(9.5,3.78)],col=INK,width=2,dashed=True)
        node(9.5,3.25,'i',TEAL,.38);node(9.5,3.78,'j',orange,.38)
        text(slide,'N',8.38,2.73,.25,.3,14,col=MUTED)
        text(slide,'C',10.98,4.0,.25,.3,14,col=MUTED)
        text(slide,'20',8.85,3.12,.45,.34,17,True,TEAL)
        text(slide,'80',8.85,3.7,.45,.34,17,True,orange)
        text(slide,info['contact_label'],10.0,3.12,1.72,.72,17,True)
        text(slide,info['left_caption'],.82,4.98,4.85,.88,18)
        text(slide,info['right_caption'],8.35,5.03,4.4,.88,18)
    else:
        assert info['kind']=='compatible_pairs'
        for index,case in enumerate(info['cases']):
            x=.8+index*4.16
            text(slide,case['title'],x,1.72,3.8,.5,22,True,TEAL)
            text(slide,'i = 20',x+.18,2.4,1.2,.4,18,col=TEAL)
            text(slide,'j = 80',x+2.13,2.4,1.2,.4,18,col=orange)
            stroke([(x+.05,3.83),(x+.65,3.65),(x+.65,3.19),(x+.27,2.93)],width=5)
            stroke([(x+3.25,2.93),(x+2.68,3.19),(x+2.68,3.65),(x+3.38,3.83)],width=5)
            stroke([(x+.65,3.19),(x+2.68,3.19)],col=INK,width=2,dashed=True)
            node(x+.65,3.19,case['residues'][0],TEAL,.8)
            node(x+2.68,3.19,case['residues'][1],orange,.8)
            text(slide,case['caption'],x,4.18,3.85,.72,20)
        text(slide,info['mechanism'],.8,5.26,11.95,.57,23,True,TEAL)
    text(slide,info['legend'],.8,6.71,11.85,.25,11,col=MUTED)


def picture(slide, filename, crop, x, y, w, h):
    path=ROOT/'assets/figures'/filename
    with Image.open(path) as im:
        iw,ih=im.size
    l,t,r,b=crop
    ratio=iw*(1-l-r)/(ih*(1-t-b))
    pw=min(w,h*ratio)
    ph=pw/ratio
    obj=slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))
    obj.crop_left,obj.crop_top,obj.crop_right,obj.crop_bottom=crop
    obj._element.nvPicPr.cNvPr.set('descr',filename+'; publisher figure, attributed in speaker notes')

sources={}
for filename in ['evolution_sources.json','coevolution_sources.json','confidence_sources.json','method_sources.json','core_sources.json']:
    incoming=json.loads((ROOT/'content'/filename).read_text())
    assert not set(incoming)&set(sources), 'duplicate source keys'
    sources.update(incoming)
data=[]
for filename in ['evolution_msa.json','coevolution.json','confidence.json','alternative_methods.json','core_workflow.json']:
    data += json.loads((ROOT/'content'/filename).read_text())
data.sort(key=lambda s:s['id'])
assert [s['id'] for s in data]==list(range(1,len(data)+1))
assert sum(s['minutes'] for s in data)==180
assert sum(s['minutes'] for s in data if s['kind']=='break')==20
assert all(s['notes'] and set(s['sources'])<=set(sources) for s in data)
prs=Presentation()
prs.slide_width,prs.slide_height=Inches(W),Inches(H)
prs.core_properties.title='MSA와 protein structure 예측: evolutionary relationship와 대안 structure 탐색'
prs.core_properties.subject='UST 2강 · 한국어 180분 · 생물학 배경 석사생'
prs.core_properties.author='Donghan Lee'
prs.core_properties.language='ko-KR'
notes=['# MSA와 protein structure 예측 · 강사용 노트','', '총 180분, 휴식 20분 포함. 각 실습 바로 다음 슬라이드에 정답·상세 해설을 제시합니다. 보충 답안은 content/worksheet_answers.md와 practice/instructor_answers.md에 있습니다.','']
elapsed=0
checks=[]
for s in data:
    slide_no=s['id']
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    kind=s['kind']
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=color(INK if kind=='break' else PAPER)
    if slide_no==1:
        text(slide,'UST · 2강',.75,.65,11.8,.45,18,True,TEAL)
        text(slide,'MSA와\nprotein structure 예측',.75,1.65,11.8,1.75,44,True)
        text(slide,'evolutionary relationship와 대안 structure 탐색',.78,3.65,11.7,.65,29,col=TEAL)
        paragraphs(slide,s['body'][:3],.8,4.65,11.7,1.95,21)
    elif kind=='break':
        text(slide,'휴식 · 10분',.8,2,11.8,1,44,True,'FFFFFF')
        text(slide,'\n'.join(s['body']),.8,3.6,11.8,1.3,24,col='FFFFFF')
    else:
        text(slide,s['section'],.68,.30,12,.35,14,col=TEAL)
        title_size=32
        if lines(s['title'],12,title_size)>1:
            title_size=30
        text(slide,s['title'],.65,.78,12,.65,title_size,True)
        if s.get('worksheet'):
            worksheet = s['worksheet']
            text(slide,worksheet['data_title'],.70,1.57,6.65,.39,19,True,TEAL)
            text(slide,'문제 · 근거와 함께 답하세요',7.65,1.57,4.95,.39,19,True,TEAL)
            if worksheet.get('table'):
                end=table(slide,worksheet['table'],2.02,17,.67,6.65,True)
                assert end <= 6.30, (slide_no,'worksheet table overlaps caption')
            else:
                code=worksheet['data_text']
                size=worksheet.get('data_size',18)
                height=(len(code.split('\n'))*size*1.13+7)/72
                assert height<=4.28,(slide_no,'worksheet data too tall')
                text(slide,code,.72,2.02,6.60,height,size,mono=worksheet.get('mono',True))
            paragraphs(slide,worksheet['questions'],7.65,2.02,4.95,4.28,20)
            text(slide,worksheet['caption'],.72,6.40,11.9,.58,17,col=TEAL)
        elif s.get('diagram'):
            coevolution_diagram(slide,s['diagram'])
            paragraphs(slide,s.get('body',[]),.8,6.0,11.9,.65,17)
        elif s.get('profile'):
            confidence_profile(slide,s['profile']['values'])
            paragraphs(slide,s.get('body',[]),.74,5.3,11.85,1.57,20)
        elif kind=='table' or s.get('table'):
            end=table(slide,s['table'])
            paragraphs(slide,s.get('body',[]),.7,end+.20,11.9,6.88-end-.2,18)
        elif kind=='sequence' or s.get('sequence'):
            code=s['sequence']['text']
            size=21
            while max(lines(row,11.9,size,True) for row in code.split('\n'))>1 and size>17:
                size-=1
            height=(len(code.split('\n'))*size*1.13+7)/72
            text(slide,code,.72,1.65,11.9,height,size,mono=True)
            y=1.65+height+.18
            explanation=s['sequence'].get('explanation',[])
            if isinstance(explanation,str): explanation=[explanation]
            explanation='\n'.join(explanation)
            eh=(lines(explanation,11.85,18)*18*1.13+4)/72
            text(slide,explanation,.74,y,11.85,eh,18,col=TEAL)
            y+=eh+.18
            paragraphs(slide,s.get('body',[]),.74,y,11.85,6.82-y,20)
        elif kind=='image':
            picture(slide,s['image'],s.get('crop',[0,0,0,0]),.65,1.55,12.0,4.10)
            text(slide,'\n'.join(s['body']),.7,5.85,11.95,.99,18)
            attribution=next(a for a in json.loads((ROOT/'assets/figures/attribution.json').read_text()) if a['file']==s['image'])
            credit=text(slide,s['image_credit'],.7,6.85,9.5,.21,9,col=MUTED)
            credit.text_frame.paragraphs[0].runs[0].hyperlink.address=attribution['article_url']
            license_link=text(slide,'CC BY 4.0',10.6,6.85,2,.21,9,col=MUTED)
            license_link.text_frame.paragraphs[0].runs[0].hyperlink.address=attribution['license_url']
        elif kind=='references':
            for i,key in enumerate(s['sources']):
                source=sources[key]
                title=source['title']
                while lines(title,11.9,18)>1 and len(title)>85:title=title[:82]+'…'
                text(slide,title,.7,1.62+i*.86,11.9,.38,18,True)
                url_label=source['url']
                if lines(url_label,11.9,12)>1:
                    url_label=url_label.split('/')[2]+' · 원문 열기'
                link=text(slide,url_label,.7,2.01+i*.86,11.9,.32,12,col=TEAL)
                for p in link.text_frame.paragraphs:
                    for run in p.runs:run.hyperlink.address=source['url']
            text(slide,'전체 출처와 논문 그림 라이선스: content/sources.json 및 발표자 노트',.7,6.35,11.9,.48,18)
        else:
            if kind=='activity':
                text(slide,'종이·구두 응답 또는 컴퓨터 실습',.74,1.72,11.85,.5,18,col=TEAL)
            paragraphs(slide,s.get('body',[]),.75,2.5 if kind=='activity' else 1.9,11.8,4.1 if kind=='activity' else 4.8,25)
    if s.get('link'):
        link=text(slide,s['link']['label'],.75,6.70,11.8,.34,18,col=TEAL)
        link.text_frame.paragraphs[0].runs[0].hyperlink.address=s['link']['url']
    col='CCD8DE' if kind=='break' else MUTED
    text(slide,f'{slide_no:02d} / {len(data)}',11.8,7.08,.95,.25,10,col=col)
    refs='\n'.join(f'[{k}] {sources[k]["title"]}\n{sources[k]["url"]}' for k in s['sources'])
    note=f'{s["title"]}\n시간: {elapsed}–{elapsed+s["minutes"]}분 ({s["minutes"]}분)\n\n{s["notes"]}\n\n출처\n{refs}'
    if s.get('image'):
        attrib=next(x for x in json.loads((ROOT/'assets/figures/attribution.json').read_text()) if x['file']==s['image'])
        note+=f'\n그림: {attrib["title"]}\n{attrib["figure_url"]}\n{attrib["license"]} {attrib["license_url"]}\n슬라이드에는 지정 패널만 확대하여 표시. 원본 파일은 보존.'
    slide.notes_slide.notes_text_frame.text=note
    notes += [f'## {slide_no}. {s["title"]} ({elapsed}–{elapsed+s["minutes"]}분)','',s['notes'],'',refs,'']
    for shape in slide.shapes:
        assert min(shape.left,shape.top)>=0
        assert shape.left+shape.width<=prs.slide_width+10 and shape.top+shape.height<=prs.slide_height+10,(slide_no,shape.name)
    checks.append({'slide':slide_no,'minutes':s['minutes'],'shapes':len(slide.shapes),'notes':bool(note)})
    elapsed+=s['minutes']
path=BUILD/'msa_structure_prediction_3h_ko.pptx'
archive(path)
prs.save(path)
(ROOT/'content/sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
(ROOT/'content/teacher_notes.md').write_text('\n'.join(notes))
(BUILD/'content_all.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
(BUILD/'deck_checks.json').write_text(json.dumps({'slides':len(data),'minutes':elapsed,'break_minutes':20,'checks':checks},ensure_ascii=False,indent=2)+'\n')
print(path)
print(f'{len(data)} slides, {elapsed} minutes, {len(sources)} sources, all bounds/notes/text estimates passed')

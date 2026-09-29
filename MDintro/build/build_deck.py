"""Build editable Korean lecture 3 (MD introduction, 60 min). Slide text lives in content/md_intro.json."""
from pathlib import Path
import json
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / 'build'
FONT = 'Arial Unicode MS'
FONT_FILE = '/Library/Fonts/Arial Unicode.ttf'
INK, TEAL, MUTED, PAPER, ORANGE = '183342', '277F87', '5A6A72', 'FAFAF6', 'BF570C'
W, H = 13.333333, 7.5
slide_no = 0


def color(value):
    return RGBColor.from_string(value)


def lines(value, width, size):
    # ponytail: font-width estimate that wraps at spaces like PowerPoint; exported pages remain the final fit check.
    font = ImageFont.truetype(FONT_FILE, round(size * 2))
    limit = (width * 72 - 6) * 2
    count = 0
    for raw in str(value).split('\n'):
        count += 1
        current = ''
        for word in raw.split(' '):
            candidate = f'{current} {word}' if current else word
            if font.getlength(candidate) <= limit:
                current = candidate
                continue
            if current:
                count += 1
            current = word
            while font.getlength(current) > limit:
                count += 1
                current = current[next(i for i in range(len(current), 0, -1) if font.getlength(current[:i]) <= limit):]
    return count


def style(run, size, bold=False, col=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    pr = run._r.get_or_add_rPr()
    pr.set('lang', 'ko-KR')
    node = OxmlElement('a:ea')
    node.set('typeface', FONT)
    pr.append(node)


def fill(frame, value, size, bold=False, col=INK, align=None):
    for i, raw in enumerate(str(value).split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.line_spacing = 1.13
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = raw
        style(run, size, bold, col)


def text(slide, value, x, y, w, h, size=22, bold=False, col=INK, align=None):
    estimate = lines(value, w, size) * size * 1.13 + 4
    assert estimate <= h * 72 + 5, (slide_no, str(value)[:80], estimate, h * 72)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    fill(frame, value, size, bold, col, align)
    return shape


def paragraphs(slide, values, x, y, w, h, size=24):
    heights = [(lines(t, w, size) * size * 1.13 + 4) / 72 for t in values]
    gap = .21
    assert sum(heights) + gap * max(0, len(heights) - 1) <= h + .04, (slide_no, 'body too tall', values)
    for value, height in zip(values, heights):
        text(slide, value, x, y, w, height, size)
        y += height + gap
    return y


def table(slide, info, y=1.55, size=19, x=.67, width=12, compact=False):
    matrix = [info['headers']] + info['rows']
    cols = len(matrix[0])
    widths = info.get('widths', [width / cols] * cols)
    assert len(widths) == cols and abs(sum(widths) - width) < .01, (slide_no, 'table widths')
    row_heights = [max(.44 if compact else .60, max(lines(str(v), widths[j] - .30, size) for j, v in enumerate(row)) * size * 1.14 / 72 + .17) for row in matrix]
    if sum(row_heights) > 4.6 and size > 17:
        return table(slide, info, y, size - 1, x, width, compact)
    assert sum(row_heights) <= 4.75, (slide_no, 'table too tall', row_heights)
    shp = slide.shapes.add_table(len(matrix), cols, Inches(x), Inches(y), Inches(width), Inches(sum(row_heights)))
    tbl = shp.table
    for column, w in zip(tbl.columns, widths):
        column.width = Inches(w)
    for i, row in enumerate(matrix):
        assert len(row) == cols, (slide_no, 'ragged table row', row)
        tbl.rows[i].height = Inches(row_heights[i])
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.055)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = color(TEAL if i == 0 else ('EEF3F2' if i % 2 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.line_spacing = 1.10
                if info.get('center') and j:
                    p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    style(run, size, i == 0, info.get('text_colors', {}).get(f'{i},{j}', 'FFFFFF' if i == 0 else INK))
    return y + sum(row_heights)


def xy_chart(slide, info, y=1.5):
    data = XyChartData()
    for s in info['series']:
        xs = info.get('x') or list(range(1, len(s['values']) + 1))
        assert len(xs) == len(s['values']), (slide_no, s['name'])
        series = data.add_series(s['name'])
        for x, v in zip(xs, s['values']):
            series.add_data_point(x, v)
    kind = XL_CHART_TYPE.XY_SCATTER_LINES if info.get('markers') else XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS
    height = info.get('height', 3.45)
    chart = slide.shapes.add_chart(kind, Inches(.7), Inches(y), Inches(12), Inches(height), data).chart
    chart.has_title = False
    chart.has_legend = len(info['series']) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.name = FONT
        chart.legend.font.size = Pt(15)
    for axis, label, scale in [(chart.category_axis, info['x_title'], info['x_scale']), (chart.value_axis, info['y_title'], info['y_scale'])]:
        axis.has_title = True
        axis.axis_title.text_frame.text = label
        for p in axis.axis_title.text_frame.paragraphs:
            for run in p.runs:
                style(run, 16)
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(14)
        axis.minimum_scale, axis.maximum_scale, axis.major_unit = scale
    chart.category_axis.has_major_gridlines = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = color('D8E1E1')
    for series, s in zip(chart.series, info['series']):
        series.format.line.color.rgb = color(s['color'])
        series.format.line.width = Pt(2.5)
        if info.get('markers'):
            series.marker.style = XL_MARKER_STYLE.CIRCLE
            series.marker.size = 6
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color(s['color'])
            series.marker.format.line.color.rgb = color(s['color'])
    return y + height


def block(slide, kind, x, y, w, h, fill_color, value='', size=19, col=INK, bold=False):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill_color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if value:
        assert lines(value, w - .2, size) * size * 1.2 / 72 <= h - .06, (slide_no, 'block text', value)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.08)
        frame.margin_top = frame.margin_bottom = Inches(.03)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill(frame, value, size, bold, col, PP_ALIGN.CENTER)
    return shape


def rule(slide, x1, y1, x2, y2, width=1.5):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color(INK)
    connector.line.width = Pt(width)


def md_loop(slide, info):
    for i, value in enumerate(info['steps']):
        x = .7 + i * 3.17
        block(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.65, 2.42, 1.6, 'F5E8D5' if i == 1 else 'DEEFEC', value, 20)
        if i < 3:
            block(slide, MSO_SHAPE.RIGHT_ARROW, x + 2.52, 2.23, .55, .44, TEAL)
    block(slide, MSO_SHAPE.LEFT_ARROW, 1.9, 3.42, 9.55, .62, TEAL, info['repeat'], 19, 'FFFFFF', True)
    return 4.25


def timescale(slide, info):
    x0, x1, lo, hi = 1.0, 12.3, -15, 0

    def pos(exponent):
        return x0 + (exponent - lo) / (hi - lo) * (x1 - x0)

    band = info['band']
    block(slide, MSO_SHAPE.RECTANGLE, x0, 1.55, pos(band['to']) - x0, 3.05, 'E3EFEE')
    text(slide, band['label'], x0 + .1, 1.6, pos(band['to']) - x0 - .2, .36, 16, True, TEAL)
    for row, bar in enumerate(info['bars']):
        y = 2.1 + row * .6
        block(slide, MSO_SHAPE.ROUNDED_RECTANGLE, pos(bar['from']), y, pos(bar['to']) - pos(bar['from']), .4, ORANGE if bar['from'] >= band['to'] else TEAL)
        text(slide, bar['label'], pos(bar['to']) + .1, y + .01, 13.1 - pos(bar['to']), .62, 17)
    axis_y = 4.75
    rule(slide, x0, axis_y, x1, axis_y)
    for exponent, label in zip(range(lo, hi + 1, 3), ['fs\n10⁻¹⁵ s', 'ps\n10⁻¹² s', 'ns\n10⁻⁹ s', 'µs\n10⁻⁶ s', 'ms\n10⁻³ s', 's\n1 s']):
        rule(slide, pos(exponent), axis_y - .08, pos(exponent), axis_y + .08)
        text(slide, label, pos(exponent) - .6, axis_y + .12, 1.2, .6, 15, align=PP_ALIGN.CENTER)
    return 5.55


data = json.loads((ROOT / 'content/md_intro.json').read_text())
sources = json.loads((ROOT / 'content/sources.json').read_text())
assert [s['id'] for s in data] == list(range(1, len(data) + 1))
assert sum(s['minutes'] for s in data) == 60
assert all(s['notes'] and set(s['sources']) <= set(sources) for s in data)
for s in data:
    if s['kind'] == 'activity':
        assert data[s['id']].get('answer_for') == s['id'], ('answer must follow its activity', s['id'])
# A term may not be shown before the slide that explains it.
explained_on = {
    'force field': 'Force field: atom 사이의 force를 정하는 근사식',
    'replica': 'MD의 기본 반복: force를 계산하고 조금 움직인다',
    'NVT': 'NVT와 NPT: 계산 중 무엇을 일정하게 두나',
    'NPT': 'NVT와 NPT: 계산 중 무엇을 일정하게 두나',
    'thermostat': 'NVT와 NPT: 계산 중 무엇을 일정하게 두나',
    'barostat': 'NVT와 NPT: 계산 중 무엇을 일정하게 두나',
    'restraints': '준비 4단계: EM → NVT → NPT → production',
    'production': '준비 4단계: EM → NVT → NPT → production',
    'RMSD': 'RMSD: 기준 structure에서 얼마나 달라졌나',
    'RMSF': 'RMSF: residue마다 얼마나 흔들렸나',
    'occupancy': '거리와 H-bond: 한 장의 값에서 시간에 따른 분포로',
    'sampling': 'MD로 답할 수 있는 것과 한계',
}
shown = [json.dumps({k: v for k, v in s.items() if k not in ('notes', 'sources')}, ensure_ascii=False) for s in data]
for term, title in explained_on.items():
    first = next(i for i, page in enumerate(shown, 1) if term in page)
    assert first >= next(s['id'] for s in data if s['title'] == title), (term, 'shown on slide', first, 'before its explanation')
# The H-bond exercise and its answer slide must agree with the stated rule.
exercise = next(s for s in data if 'hbond_example' in s)
hb = exercise['hbond_example']
near = [d <= hb['max_distance'] for d in hb['distance']]
angle_ok = [a <= hb['max_angle'] for a in hb['angle']]
both = [n and a for n, a in zip(near, angle_ok)]
assert hb['max_distance'] not in hb['distance'] and hb['max_angle'] not in hb['angle'], 'avoid boundary values'
assert (sum(near), sum(both)) == (hb['distance_only'], hb['both']) == (8, 7)
shown = exercise['worksheet']['table']['rows']
assert shown[0][1:] == [f'{d:.1f}' for d in hb['distance']] and shown[1][1:] == [str(a) for a in hb['angle']]
marks = data[exercise['id']]['table']['rows']
assert [row[1:] for row in marks] == [['O' if ok else 'X' for ok in flags] for flags in (near, angle_ok, both)]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Molecular dynamics(MD) 입문'
prs.core_properties.subject = 'UST 3강 · 한국어 60분 개념 강의 · 이어서 2시간 실습'
prs.core_properties.author = 'Donghan Lee'
prs.core_properties.language = 'ko-KR'
notes = ['# MD 입문 · 강사용 노트', '', '총 60분(휴식 없음). 각 실습 바로 다음 슬라이드에 정답·해설이 있으며, 이어서 2시간 실습을 진행합니다.', '']
elapsed = 0
for s in data:
    slide_no = s['id']
    kind = s['kind']
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(PAPER)
    if kind == 'cover':
        text(slide, 'UST · 3강', .75, .65, 11.8, .45, 18, True, TEAL)
        text(slide, s['title'], .75, 1.65, 11.8, 1.75, 44, True)
        text(slide, s['subtitle'], .78, 3.65, 11.7, .65, 29, col=TEAL)
        paragraphs(slide, s['body'], .8, 4.65, 11.7, 1.95, 21)
    else:
        text(slide, s['section'], .68, .30, 12, .35, 14, col=TEAL)
        title_size = 32 if lines(s['title'], 12, 32) == 1 else 30
        text(slide, s['title'], .65, .78, 12, .65, title_size, True)
        if kind == 'activity' and s.get('worksheet'):
            sheet = s['worksheet']
            text(slide, sheet['data_title'], .70, 1.52, 11.9, .39, 19, True, TEAL)
            end = table(slide, sheet['table'], 1.98, 18, .67, 12, True)
            text(slide, '문제 · 근거와 함께 답하세요', .70, end + .22, 11.9, .39, 19, True, TEAL)
            paragraphs(slide, sheet['questions'], .75, end + .68, 11.8, 6.45 - end - .68, 20)
            if sheet.get('caption'):
                text(slide, sheet['caption'], .72, 6.55, 11.9, .4, 16, col=TEAL)
        elif kind == 'activity':
            text(slide, '종이 · 구두 응답', .74, 1.72, 11.85, .5, 18, col=TEAL)
            paragraphs(slide, s['body'], .75, 2.5, 11.8, 4.1, 25)
        elif kind == 'table':
            end = table(slide, s['table'])
            paragraphs(slide, s.get('body', []), .7, end + .20, 11.9, 6.88 - end - .2, 18)
        elif kind in ('chart', 'diagram'):
            if kind == 'chart':
                end = xy_chart(slide, s['chart'])
            else:
                end = {'md_loop': md_loop, 'timescale': timescale}[s['diagram']['kind']](slide, s['diagram'])
            bottom = 6.55 if s.get('caption') else 6.88
            paragraphs(slide, s['body'], .74, end + .18, 11.85, bottom - end - .18, 18 if kind == 'chart' else 20)
            if s.get('caption'):
                text(slide, s['caption'], .74, 6.62, 11.85, .3, 12, col=MUTED)
        elif kind == 'text':
            paragraphs(slide, s['body'], .75, 1.9, 11.8, 4.8, 25)
        elif kind == 'references':
            for i, key in enumerate(s['sources']):
                y = 1.5 + i * .44
                text(slide, sources[key]['title'], .7, y, 11.9, .27, 14, True)
                link = text(slide, sources[key]['url'], .7, y + .24, 11.9, .2, 10, col=TEAL)
                link.text_frame.paragraphs[0].runs[0].hyperlink.address = sources[key]['url']
        else:
            raise ValueError((slide_no, kind))
    text(slide, f'{slide_no:02d} / {len(data)}', 11.8, 7.08, .95, .25, 10, col=MUTED)
    refs = '\n'.join(f'[{k}] {sources[k]["title"]}\n{sources[k]["url"]}' for k in s['sources'])
    slide.notes_slide.notes_text_frame.text = f'{s["title"]}\n시간: {elapsed}–{elapsed + s["minutes"]}분 ({s["minutes"]}분)\n\n{s["notes"]}\n\n출처\n{refs}'
    notes += [f'## {slide_no}. {s["title"]} ({elapsed}–{elapsed + s["minutes"]}분)', '', s['notes'], '', refs, '']
    for shape in slide.shapes:
        assert min(shape.left, shape.top) >= 0, (slide_no, shape.name)
        assert shape.left + shape.width <= prs.slide_width + 10 and shape.top + shape.height <= prs.slide_height + 10, (slide_no, shape.name)
    elapsed += s['minutes']

path = BUILD / 'md_intro_1h_ko.pptx'
prs.save(path)
(ROOT / 'content/teacher_notes.md').write_text('\n'.join(notes))
print(path)
print(f'{len(data)} slides, {elapsed} minutes, {len(sources)} sources, bounds/notes/text estimates passed')

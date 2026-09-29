"""Check timings, editable deck, rendered text, worksheet, and protected lesson 1."""
from pathlib import Path
import hashlib
import json
import re
import unicodedata
import pymupdf as fitz
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
build=ROOT/'build'
data=json.loads((build/'content_all.json').read_text())
assert len(data)==104 and [s['id'] for s in data]==list(range(1,105))
assert sum(s['minutes'] for s in data)==180
assert sum(s['minutes'] for s in data if s['kind']=='break')==20
prs=Presentation(build/'msa_structure_prediction_3h_ko.pptx')
pdf=fitz.open(build/'msa_structure_prediction_3h_ko.pdf')
worksheet=fitz.open(build/'msa_structure_prediction_worksheet_ko.pdf')
assert len(prs.slides)==len(pdf)==104 and len(worksheet)==7
activities=[s for s in data if s['kind']=='activity' or s['title'].startswith('바로 실습')]
assert len(activities)==16
for activity in activities:
    answer=data[activity['id']]
    assert answer.get('answer_for')==activity['id'], ('answer not immediately after exercise', activity['id'])
answers=[s for s in data if 'answer_for' in s]
assert len(answers)==24
expected_worksheet_items={f'{page}.{q}' for page, count in [(1,3),(2,4),(3,4),(4,4),(5,4),(6,4),(7,3)] for q in range(1,count+1)} | {'5.table'}
worksheet_items=[item for slide in data for item in slide.get('worksheet_items',[])]
assert len(worksheet_items)==len(set(worksheet_items))==27
assert set(worksheet_items)==expected_worksheet_items
assert sum(bool(s.get('worksheet')) for s in data)==11
fixture_sequences=[line for line in (ROOT/'practice/fictional_learning.a3m').read_text().splitlines() if line and not line.startswith('>')]
a3m_question=next(s for s in data if '2.1' in s.get('worksheet_items',[]))
embedded_sequences=[line for line in a3m_question['worksheet']['data_text'].splitlines() if line and not line.startswith('>')]
assert embedded_sequences==fixture_sequences
candidate_question=next(s for s in data if '7.1' in s.get('worksheet_items',[]))
fixture_candidates=json.loads((ROOT/'practice/synthetic_comparison.json').read_text())['rows']
assert len(candidate_question['worksheet']['table']['rows'])==len(fixture_candidates)==8
for actual,expected in zip(candidate_question['worksheet']['table']['rows'],fixture_candidates):
    assert actual[0]==expected['id']
    assert [float(v) for v in actual[2:]]==[expected['mean_plddt'],expected['distance_A'],expected['inter_region_pae_A']]
for answer in answers:
    previous=data[answer['id']-2]
    assert previous['id']==answer['answer_for'] or previous.get('answer_for')==answer['answer_for']
    assert answer['body'] and answer['notes']
    assert answer.get('table') or answer.get('sequence')
assert not re.search(r'\b(?:boltz|chai|esmfold|esm-fold)\b',json.dumps(data).lower())
assert not re.search(r'공진화|공변이|공변화|종분화|종 분화|유전자|단백질|아미노산|잔기|도메인|신뢰도|마스킹|클러스터',json.dumps(data,ensure_ascii=False))
assert not any(word in json.dumps(data,ensure_ascii=False) for word in ['공evolution','gene redundancy','phylogeny수','연binding','coordinates틀','파일 읽기 training용','speciation·redundancy'])

def norm(s):
    return ''.join(unicodedata.normalize('NFKC',s).split()).replace('\u200b','')

missing=[]
for i,(slide,page) in enumerate(zip(prs.slides,pdf),1):
    assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text
    rendered=norm(page.get_text())
    assert '\ufffd' not in rendered
    source=data[i-1]
    if source.get('diagram'):
        assert any(shp._element.tag.endswith('cxnSp') for shp in slide.shapes), (i,'diagram connectors missing')
        assert not any(shp._element.tag.endswith('pic') for shp in slide.shapes), (i,'diagram should remain editable')
        assert norm(source['diagram']['legend']) in rendered
    if source.get('link'):
        assert source['link']['url'] in [item.get('uri') for item in page.get_links()], (i,'missing PDF hyperlink')
        assert source['link']['label']==source['link']['url'], (i,'URL hidden by label')
        assert norm(source['link']['url']) in rendered, (i,'full URL missing from slide')
    if source.get('profile'):
        charts=[shp.chart for shp in slide.shapes if shp.has_chart]
        assert len(charts)==1
        assert tuple(charts[0].series[0].values)==tuple(source['profile']['values'])
        assert charts[0].value_axis.minimum_scale==0 and charts[0].value_axis.maximum_scale==100
        assert 'pLDDT' in rendered and norm('residue number') in rendered
    if source.get('lddt_example'):
        example=source['lddt_example']
        assert all(d < example['inclusion_radius'] for d in example['reference_distances'])
        errors=[abs(p-r) for p,r in zip(example['predicted_distances'],example['reference_distances'])]
        passes=[sum(error < tolerance for tolerance in example['thresholds']) for error in errors]
        assert passes==example['passes']
        assert sum(passes)/(len(errors)*len(example['thresholds']))==example['score']==0.625
        changed_passes=sum(abs(19-example['reference_distances'][-1]) < t for t in example['thresholds'])
        assert (sum(passes[:-1])+changed_passes)/16==0.5625
    if source.get('plddt_example'):
        example=source['plddt_example']
        assert abs(sum(example['probabilities'])-1)<1e-9
        assert all(any(abs(c-(0.01+0.02*k))<1e-9 for k in range(50)) for c in example['bin_centers'])
        expected_score=100*sum(p*c for p,c in zip(example['probabilities'],example['bin_centers']))
        assert abs(expected_score-example['expected_plddt'])<1e-9
    expected=[] if source['kind']=='references' else source.get('body',[])[:]
    if source.get('worksheet'):
        worksheet_source=source['worksheet']
        expected += [worksheet_source['data_title'],worksheet_source['caption']] + worksheet_source['questions']
        if worksheet_source.get('table'):
            expected += worksheet_source['table']['headers'] + [str(v) for row in worksheet_source['table']['rows'] for v in row]
        if worksheet_source.get('data_text'):
            expected += worksheet_source['data_text'].splitlines()
    if source.get('table'):
        expected += source['table']['headers'] + [str(v) for row in source['table']['rows'] for v in row]
    if source.get('sequence'):
        expected += source['sequence']['text'].splitlines()
    for value in expected:
        assert not norm(value) or norm(value) in rendered, (i,'source content missing',value)
    for shp in slide.shapes:
        assert shp.left>=0 and shp.top>=0
        assert shp.left+shp.width<=prs.slide_width+10
        assert shp.top+shp.height<=prs.slide_height+10
        chunks=[]
        if shp.has_text_frame:
            chunks=[p.text for p in shp.text_frame.paragraphs]
            if shp.top > prs.slide_height * .92:
                assert not re.search(r'\d+\s*[–-]\s*\d+\s*분', ' '.join(chunks)), (i,'timed footer')
        if shp.has_table:chunks=[c.text for row in shp.table.rows for c in row.cells]
        for chunk in chunks:
            if norm(chunk) and norm(chunk) not in rendered:missing.append({'slide':i,'text':chunk})
    for block in page.get_text('blocks'):
        assert block[0]>=-1 and block[1]>=-1 and block[2]<=page.rect.width+1 and block[3]<=page.rect.height+1,(i,block[:4])
for i,page in enumerate(worksheet,1):
    assert '\ufffd' not in page.get_text()
    for block in page.get_text('blocks'):
        assert block[0]>=-1 and block[1]>=-1 and block[2]<=page.rect.width+1 and block[3]<=page.rect.height+1,('worksheet',i,block[:4])
original=json.loads((build/'lecture1_baseline_sha256.json').read_text())
for path,digest in original.items():
    assert hashlib.sha256((ROOT.parent/path).read_bytes()).hexdigest()==digest,path
query=''.join((ROOT/'examples/calmodulin_P0DP23.fasta').read_text().splitlines()[1:])
assert len(query)==149
report={'slides':104,'lecture_pdf_pages':104,'worksheet_pages':7,'total_minutes':180,'break_minutes':20,'all_slide_notes_present':True,'lecture1_outputs_unchanged':True,'missing_rendered_text':missing,'exercises_with_immediate_answers':16,'answer_and_execution_check_slides':24,'worksheet_items_embedded':27,'actual_structure_inference_run':False,'synthetic_practice_data':True}
(build/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
assert not missing,missing
print(json.dumps(report,ensure_ascii=False,indent=2))

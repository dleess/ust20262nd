"""Build editable Korean lecture 2. Source content lives in content/*.json."""
from pathlib import Path
from datetime import datetime
import json
import shutil
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / 'build'
FONT = 'Arial Unicode MS'
FONT_FILE = '/Library/Fonts/Arial Unicode.ttf'
MONO_FILE = '/System/Library/Fonts/Supplemental/Courier New.ttf'
INK, TEAL, MUTED, PAPER = '183342', '277F87', '5A6A72', 'FAFAF6'
W, H = 13.333333, 7.5

def archive(path):
    if path.exists():
        folder = BUILD / 'archive' / datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        folder.mkdir(parents=True)
        shutil.copy2(path, folder / path.name)

def color(value):
    return RGBColor.from_string(value)

def lines(value, width, size, mono=False):
    # ponytail: font-width estimate; exported PowerPoint pages remain the final fit check.
    font = ImageFont.truetype(MONO_FILE if mono else FONT_FILE, round(size * 2))
    limit = (width * 72 - 6) * 2
    count = 0
    for raw in str(value).split('\n'):
        count += 1
        current = ''
        for char in raw:
            if font.getlength(current + char) > limit:
                count += 1
                current = char
            else:
                current += char
    return count

def style(run, size, bold=False, col=INK, mono=False):
    run.font.name = 'Courier New' if mono else FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color(col)
    pr = run._r.get_or_add_rPr()
    pr.set('lang', 'ko-KR')
    node = OxmlElement('a:ea')
    node.set('typeface', FONT)
    pr.append(node)

def text(slide, value, x, y, w, h, size=22, bold=False, col=INK, mono=False):
    estimate = lines(value, w, size, mono) * size * 1.13 + 4
    assert estimate <= h * 72 + 5, (slide_no, str(value)[:80], estimate, h * 72)
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = not mono
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    for i, raw in enumerate(str(value).split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.line_spacing = Pt(size * 1.13) if mono else 1.13
        r = p.add_run()
        r.text = raw
        style(r, size, bold, col, mono)
    return shape

def paragraphs(slide, values, x, y, w, h, size=24):
    if isinstance(values, str):
        values = [values]
    heights = [(lines(t, w, size) * size * 1.13 + 4) / 72 for t in values]
    gap = .21
    assert sum(heights) + gap * max(0, len(heights)-1) <= h + .04, (slide_no, 'body too tall',values)
    for val, height in zip(values, heights):
        text(slide, val, x, y, w, height, size)
        y += height + gap
    return y

def table(slide, info, y=1.55, size=19):
    matrix = [info['headers']] + info['rows']
    cols = len(matrix[0])
    widths = info.get('widths', [12.0 / cols] * cols)
    assert len(widths) == cols and abs(sum(widths)-12)<.01
    row_heights = [max(.60, max(lines(str(v), widths[j]-.30,size) for j,v in enumerate(row))*size*1.14/72+.17) for row in matrix]
    if sum(row_heights)>4.6 and size>17:
        return table(slide,info,y,size-1)
    assert sum(row_heights)<=4.75, (slide_no,'table too tall',row_heights)
    shp = slide.shapes.add_table(len(matrix),cols,Inches(.67),Inches(y),Inches(12),Inches(sum(row_heights)))
    tbl = shp.table
    for c, width in zip(tbl.columns,widths):
        c.width = Inches(width)
    for i, row in enumerate(matrix):
        assert len(row)==cols
        tbl.rows[i].height=Inches(row_heights[i])
        for j, value in enumerate(row):
            cell=tbl.cell(i,j)
            cell.text=str(value)
            cell.margin_left=cell.margin_right=Inches(.12)
            cell.margin_top=cell.margin_bottom=Inches(.055)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb=color(TEAL if i==0 else ('EEF3F2' if i%2 else 'FFFFFF'))
            for p in cell.text_frame.paragraphs:
                p.line_spacing=1.10
                for run in p.runs:
                    style(run,size,i==0,'FFFFFF' if i==0 else INK)
    return y+sum(row_heights)

def picture(slide, filename, crop, x, y, w, h):
    path=ROOT/'assets/figures'/filename
    with Image.open(path) as im:
        iw,ih=im.size
    l,t,r,b=crop
    ratio=iw*(1-l-r)/(ih*(1-t-b))
    pw=min(w,h*ratio)
    ph=pw/ratio
    obj=slide.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),Inches(pw),Inches(ph))
    obj.crop_left,obj.crop_top,obj.crop_right,obj.crop_bottom=crop
    obj._element.nvPicPr.cNvPr.set('descr',filename+'; publisher figure, attributed in speaker notes')

sources={}
for filename in ['evolution_sources.json','coevolution_sources.json','method_sources.json','core_sources.json']:
    incoming=json.loads((ROOT/'content'/filename).read_text())
    assert not set(incoming)&set(sources), 'duplicate source keys'
    sources.update(incoming)
data=[]
for filename in ['evolution_msa.json','coevolution.json','alternative_methods.json','core_workflow.json']:
    data += json.loads((ROOT/'content'/filename).read_text())
data.sort(key=lambda s:s['id'])
assert [s['id'] for s in data]==list(range(1,len(data)+1))
assert sum(s['minutes'] for s in data)==180
assert sum(s['minutes'] for s in data if s['kind']=='break')==20
assert all(s['notes'] and set(s['sources'])<=set(sources) for s in data)
prs=Presentation()
prs.slide_width,prs.slide_height=Inches(W),Inches(H)
prs.core_properties.title='MSA와 단백질 구조 예측: 진화 관계와 대안 구조 탐색'
prs.core_properties.subject='UST 2강 · 한국어 180분 · 생물학 배경 석사생'
prs.core_properties.author='Donghan Lee'
prs.core_properties.language='ko-KR'
notes=['# MSA와 단백질 구조 예측 · 강사용 노트','', '총 180분, 휴식 20분 포함. 활동 답안은 이 강사용 노트와 practice/instructor_answers.md에 수록합니다.','']
elapsed=0
checks=[]
for s in data:
    slide_no=s['id']
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    kind=s['kind']
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=color(INK if kind=='break' else PAPER)
    if slide_no==1:
        text(slide,'UST · 2강',.75,.65,11.8,.45,18,True,TEAL)
        text(slide,'MSA와\n단백질 구조 예측',.75,1.65,11.8,1.75,44,True)
        text(slide,'진화 관계와 대안 구조 탐색',.78,3.65,11.7,.65,29,col=TEAL)
        paragraphs(slide,s['body'][:3],.8,4.65,11.7,1.95,21)
    elif kind=='break':
        text(slide,'휴식 · 10분',.8,2,11.8,1,44,True,'FFFFFF')
        text(slide,'\n'.join(s['body']),.8,3.6,11.8,1.3,24,col='FFFFFF')
    else:
        text(slide,s['section'],.68,.30,12,.35,14,col=TEAL)
        title_size=32
        if lines(s['title'],12,title_size)>1:
            title_size=30
        text(slide,s['title'],.65,.78,12,.65,title_size,True)
        if kind=='table' or s.get('table'):
            end=table(slide,s['table'])
            paragraphs(slide,s.get('body',[]),.7,end+.20,11.9,6.88-end-.2,18)
        elif kind=='sequence' or s.get('sequence'):
            code=s['sequence']['text']
            size=21
            while max(lines(row,11.9,size,True) for row in code.split('\n'))>1 and size>17:
                size-=1
            height=(len(code.split('\n'))*size*1.13+7)/72
            text(slide,code,.72,1.65,11.9,height,size,mono=True)
            y=1.65+height+.18
            explanation=s['sequence'].get('explanation',[])
            if isinstance(explanation,str): explanation=[explanation]
            explanation='\n'.join(explanation)
            eh=(lines(explanation,11.85,18)*18*1.13+4)/72
            text(slide,explanation,.74,y,11.85,eh,18,col=TEAL)
            y+=eh+.18
            paragraphs(slide,s.get('body',[]),.74,y,11.85,6.82-y,20)
        elif kind=='image':
            picture(slide,s['image'],s.get('crop',[0,0,0,0]),.65,1.55,12.0,4.10)
            text(slide,'\n'.join(s['body']),.7,5.85,11.95,.99,18)
            attribution=next(a for a in json.loads((ROOT/'assets/figures/attribution.json').read_text()) if a['file']==s['image'])
            credit=text(slide,s['image_credit'],.7,6.85,9.5,.21,9,col=MUTED)
            credit.text_frame.paragraphs[0].runs[0].hyperlink.address=attribution['article_url']
            license_link=text(slide,'CC BY 4.0',10.6,6.85,2,.21,9,col=MUTED)
            license_link.text_frame.paragraphs[0].runs[0].hyperlink.address=attribution['license_url']
        elif kind=='references':
            for i,key in enumerate(s['sources']):
                source=sources[key]
                title=source['title']
                while lines(title,11.9,18)>1 and len(title)>85:title=title[:82]+'…'
                text(slide,title,.7,1.62+i*.86,11.9,.38,18,True)
                link=text(slide,source['url'],.7,2.01+i*.86,11.9,.32,12,col=TEAL)
                for p in link.text_frame.paragraphs:
                    for run in p.runs:run.hyperlink.address=source['url']
            text(slide,'전체 출처와 논문 그림 라이선스: content/sources.json 및 발표자 노트',.7,6.35,11.9,.48,18)
        else:
            if kind=='activity':
                text(slide,'종이·구두 응답 또는 컴퓨터 실습',.74,1.72,11.85,.5,18,col=TEAL)
            paragraphs(slide,s.get('body',[]),.75,2.5 if kind=='activity' else 1.9,11.8,4.1 if kind=='activity' else 4.8,25)
    if s.get('link'):
        link=text(slide,s['link']['label'],.75,6.70,11.8,.34,18,col=TEAL)
        link.text_frame.paragraphs[0].runs[0].hyperlink.address=s['link']['url']
    col='CCD8DE' if kind=='break' else MUTED
    text(slide,f'{elapsed:03d}–{elapsed+s["minutes"]:03d}분',.7,7.08,4,.25,10,col=col)
    text(slide,f'{slide_no:02d} / {len(data)}',11.8,7.08,.95,.25,10,col=col)
    refs='\n'.join(f'[{k}] {sources[k]["title"]}\n{sources[k]["url"]}' for k in s['sources'])
    note=f'{s["title"]}\n시간: {elapsed}–{elapsed+s["minutes"]}분 ({s["minutes"]}분)\n\n{s["notes"]}\n\n출처\n{refs}'
    if s.get('image'):
        attrib=next(x for x in json.loads((ROOT/'assets/figures/attribution.json').read_text()) if x['file']==s['image'])
        note+=f'\n그림: {attrib["title"]}\n{attrib["figure_url"]}\n{attrib["license"]} {attrib["license_url"]}\n슬라이드에는 지정 패널만 확대하여 표시. 원본 파일은 보존.'
    slide.notes_slide.notes_text_frame.text=note
    notes += [f'## {slide_no}. {s["title"]} ({elapsed}–{elapsed+s["minutes"]}분)','',s['notes'],'',refs,'']
    for shape in slide.shapes:
        assert min(shape.left,shape.top)>=0
        assert shape.left+shape.width<=prs.slide_width+10 and shape.top+shape.height<=prs.slide_height+10,(slide_no,shape.name)
    checks.append({'slide':slide_no,'minutes':s['minutes'],'shapes':len(slide.shapes),'notes':bool(note)})
    elapsed+=s['minutes']
path=BUILD/'msa_structure_prediction_3h_ko.pptx'
archive(path)
prs.save(path)
(ROOT/'content/sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2)+'\n')
(ROOT/'content/teacher_notes.md').write_text('\n'.join(notes))
(BUILD/'content_all.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
(BUILD/'deck_checks.json').write_text(json.dumps({'slides':len(data),'minutes':elapsed,'break_minutes':20,'checks':checks},ensure_ascii=False,indent=2)+'\n')
print(path)
print(f'{len(data)} slides, {elapsed} minutes, {len(sources)} sources, all bounds/notes/text estimates passed')
